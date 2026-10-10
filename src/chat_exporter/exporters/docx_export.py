"""Turns a Chat into a Word (.docx) file using python-docx.

Unlike the PDF export, Word files are not built from HTML. We create
the document piece by piece: a title, then for each message a bold
label ("You" / "AI") followed by its paragraphs.
"""

from pathlib import Path  # cross-platform file paths

from docx import Document  # python-docx: creates and edits Word documents

from chat_exporter.exporters.html_builder import role_label  # reuse "You" / "AI"
from chat_exporter.models import Chat  # our shared chat structure


def export_docx(chat: Chat, output_path: Path) -> Path:
    """Write the chat to a .docx file and return the path of that file."""
    document = Document()  # a new, empty Word document

    # Add the chat title as a heading (level 1 = biggest heading style).
    document.add_heading(chat.title, level=1)

    for message in chat.messages:
        # Label line: "You" or "AI" in bold.
        label_paragraph = document.add_paragraph()
        label_run = label_paragraph.add_run(role_label(message.role))
        label_run.bold = True

        # One Word paragraph for every blank-line-separated block of text.
        for block in message.text.split("\n\n"):
            if block.strip():  # skip empty blocks
                document.add_paragraph(block.strip())

    document.save(str(output_path))  # write the file to disk
    return output_path